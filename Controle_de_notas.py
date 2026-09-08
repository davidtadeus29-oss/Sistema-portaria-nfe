# -*- coding: utf-8 -*-
"""
===============================================================================
SISTEMA DE PORTARIA: BIPAGEM DE NF-e, CONTROLE DE DESVIOS E JUSTIFICATIVAS
LEITURA 100% AUTOMÁTICA VIA LEITOR DE CÓDIGO DE BARRAS (ENTER NATIVO)
===============================================================================
"""

import os
import sys
import re
import json
import sqlite3
import smtplib
import subprocess
import threading
import time
from datetime import datetime
from email.mime.text import MIMEText
from email.mime.multipart import MIMEMultipart
import tkinter as tk
from tkinter import ttk, messagebox, filedialog

# Suporte ao envio nativo pelo Outlook Desktop
try:
    import win32com.client
    import pythoncom
    OUTLOOK_COM_DISPONIVEL = True
except ImportError:
    OUTLOOK_COM_DISPONIVEL = False

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
except ImportError:
    openpyxl = None

try:
    import schedule
except ImportError:
    schedule = None


# =============================================================================
# DIRETÓRIO ABSOLUTO E BANCO DE DADOS
# =============================================================================

if getattr(sys, 'frozen', False):
    BASE_DIR = os.path.dirname(sys.executable)
else:
    BASE_DIR = os.path.dirname(os.path.abspath(__file__)) if '__file__' in globals() else os.getcwd()

DB_NAME = os.path.join(BASE_DIR, "bipagem_nfe.db")
CONFIG_FILE = os.path.join(BASE_DIR, "config_bipagem.json")

CONFIG_PADRAO = {
    "metodo_envio": "OUTLOOK_APP",  # "OUTLOOK_APP" ou "SMTP"
    "exibir_janela_email": False,
    "smtp_server": "smtp.office365.com",
    "smtp_port": 587,
    "email_remetente": "",
    "senha_app": "",
    "destinatarios": [],
    "horario_resumo": "22:00"
}


def obter_conexao():
    conn = sqlite3.connect(DB_NAME)
    conn.row_factory = sqlite3.Row
    return conn


def carregar_config():
    if os.path.exists(CONFIG_FILE):
        try:
            with open(CONFIG_FILE, "r", encoding="utf-8") as f:
                return {**CONFIG_PADRAO, **json.load(f)}
        except Exception:
            return CONFIG_PADRAO
    else:
        salvar_config(CONFIG_PADRAO)
        return CONFIG_PADRAO


def salvar_config(dados):
    try:
        with open(CONFIG_FILE, "w", encoding="utf-8") as f:
            json.dump(dados, f, indent=4, ensure_ascii=False)
    except Exception as e:
        print(f"Erro ao salvar config: {e}")


def inicializar_banco():
    conn = obter_conexao()
    cursor = conn.cursor()
    cursor.execute("""
        CREATE TABLE IF NOT EXISTS notas (
            chave TEXT PRIMARY KEY,
            numero_nf TEXT,
            serie TEXT,
            data_bip1 TEXT,
            hora_bip1 TEXT,
            dt_completa_bip1 TEXT,
            data_bip2 TEXT,
            hora_bip2 TEXT,
            dt_completa_bip2 TEXT,
            tempo_decorrido TEXT,
            minutos_decorridos REAL,
            status TEXT,
            justificativa TEXT,
            qtd_bipagens INTEGER DEFAULT 1,
            observacao TEXT
        )
    """)
    conn.commit()

    cursor.execute("PRAGMA table_info(notas)")
    colunas_existentes = [info["name"] for info in cursor.fetchall()]

    colunas_necessarias = {
        "justificativa": "TEXT",
        "qtd_bipagens": "INTEGER DEFAULT 1",
        "observacao": "TEXT",
        "dt_completa_bip1": "TEXT",
        "dt_completa_bip2": "TEXT",
        "tempo_decorrido": "TEXT",
        "minutos_decorridos": "REAL",
        "status": "TEXT"
    }

    for col_nome, col_tipo in colunas_necessarias.items():
        if col_nome not in colunas_existentes:
            try:
                cursor.execute(f"ALTER TABLE notas ADD COLUMN {col_nome} {col_tipo}")
            except Exception:
                pass

    conn.commit()
    conn.close()


def sanitizar_e_extrair_chave(texto_bruto: str):
    if not texto_bruto:
        return {"valida": False, "chave": "", "numero_nf": "N/D", "serie": "N/D"}

    apenas_numeros = re.sub(r"\D", "", texto_bruto)

    if len(apenas_numeros) >= 44:
        chave = apenas_numeros[:44]
        try:
            serie = str(int(chave[22:25]))
            numero_nf = str(int(chave[25:34]))
        except Exception:
            serie = chave[22:25]
            numero_nf = chave[25:34]
            
        return {"valida": True, "chave": chave, "numero_nf": numero_nf, "serie": serie}

    return {"valida": False, "chave": apenas_numeros, "numero_nf": "N/D", "serie": "N/D"}


def calcular_diferenca_tempo(dt_inicio: datetime, dt_fim: datetime):
    delta = dt_fim - dt_inicio
    total_segundos = max(0, int(delta.total_seconds()))
    dias, resto = divmod(total_segundos, 86400)
    horas, resto = divmod(resto, 3600)
    minutos, segundos = divmod(resto, 60)
    
    partes = []
    if dias > 0:
        partes.append(f"{dias}d")
    if horas > 0 or dias > 0:
        partes.append(f"{horas}h")
    partes.append(f"{minutos}m {segundos}s")
    
    texto_formatado = " ".join(partes)
    minutos_totais = round(total_segundos / 60.0, 2)
    return texto_formatado, minutos_totais


# =============================================================================
# MÓDULO DE ENVIO DE E-MAILS (OUTLOOK DESKTOP / POWERSHELL / SMTP)
# =============================================================================

def enviar_via_powershell_outlook(assunto: str, corpo_html: str, destinatarios: list, abrir_janela=False):
    try:
        dest_str = "; ".join(destinatarios)
        temp_html = os.path.join(BASE_DIR, "temp_email_body.html")
        with open(temp_html, "w", encoding="utf-8") as f:
            f.write(corpo_html)

        acao_envio = "$mail.Display()" if abrir_janela else "$mail.Send(); $ns.SendAndReceive($false)"

        ps_script = f"""
        $ol = New-Object -ComObject Outlook.Application
        $ns = $ol.GetNamespace('MAPI')
        $ns.Logon('', '', $false, $false)
        $mail = $ol.CreateItem(0)
        $mail.Subject = '{assunto}'
        $mail.To = '{dest_str}'
        $htmlContent = Get-Content -Path '{temp_html}' -Raw -Encoding UTF8
        $mail.HTMLBody = $htmlContent
        {acao_envio}
        """

        process = subprocess.run(
            ["powershell", "-NoProfile", "-ExecutionPolicy", "Bypass", "-Command", ps_script],
            capture_output=True,
            text=True,
            timeout=25
        )

        if os.path.exists(temp_html):
            os.remove(temp_html)

        if process.returncode == 0:
            print(f"[{datetime.now().strftime('%H:%M:%S')}] [OK] E-mail enviado via PowerShell Outlook!")
            return True, "E-mail enviado via Outlook com sucesso!"
        else:
            return False, f"Falha no PowerShell: {process.stderr}"
    except Exception as e:
        return False, f"Erro ao acionar PowerShell: {str(e)}"


def enviar_email_outlook_com(assunto: str, corpo_html: str, destinatarios: list, abrir_janela=False):
    if not OUTLOOK_COM_DISPONIVEL:
        return enviar_via_powershell_outlook(assunto, corpo_html, destinatarios, abrir_janela)

    try:
        pythoncom.CoInitialize()
        outlook = win32com.client.Dispatch("Outlook.Application")
        ns = outlook.GetNamespace("MAPI")
        ns.Logon("", "", False, False)

        mail = outlook.CreateItem(0)
        mail.Subject = assunto

        for dest in destinatarios:
            if dest.strip():
                r = mail.Recipients.Add(dest.strip())
                r.Type = 1
        mail.Recipients.ResolveAll()

        mail.HTMLBody = corpo_html

        if abrir_janela:
            mail.Display()
        else:
            mail.Save()
            mail.Send()
            try:
                ns.SendAndReceive(False)
            except Exception:
                pass

        pythoncom.CoUninitialize()
        print(f"[{datetime.now().strftime('%H:%M:%S')}] [OK] E-mail enviado e registrado no Outlook!")
        return True, "E-mail enviado via Outlook Desktop com sucesso!"
    except Exception as e:
        pythoncom.CoUninitialize()
        return enviar_via_powershell_outlook(assunto, corpo_html, destinatarios, abrir_janela)


def enviar_email_smtp_direto(assunto: str, corpo_html: str, destinatarios: list):
    config = carregar_config()
    remetente = config.get("email_remetente", "").strip()
    senha = config.get("senha_app", "").strip()
    smtp_server = config.get("smtp_server", "smtp.office365.com").strip()
    smtp_port = int(config.get("smtp_port", 587))
    
    if not remetente or not senha:
        return False, "E-mail ou senha de aplicativo não configurados."

    try:
        msg = MIMEMultipart("alternative")
        msg["Subject"] = assunto
        msg["From"] = remetente
        msg["To"] = ", ".join(destinatarios)
        msg.attach(MIMEText(corpo_html, "html", "utf-8"))
        
        with smtplib.SMTP(smtp_server, smtp_port, timeout=20) as server:
            server.starttls()
            server.login(remetente, senha)
            server.sendmail(remetente, destinatarios, msg.as_string())
            
        print(f"[{datetime.now().strftime('%H:%M:%S')}] [OK] E-mail enviado via SMTP!")
        return True, "E-mail enviado via SMTP com sucesso!"
    except Exception as e:
        print(f"[{datetime.now().strftime('%H:%M:%S')}] [ERRO SMTP] {e}")
        return False, f"Falha no SMTP: {str(e)}"


def disparar_email_geral(assunto: str, corpo_html: str, destinatarios: list):
    if not destinatarios:
        return False, "Nenhum destinatário cadastrado na lista."

    config = carregar_config()
    metodo = config.get("metodo_envio", "OUTLOOK_APP")
    abrir_janela = config.get("exibir_janela_email", False)

    if metodo == "OUTLOOK_APP":
        return enviar_email_outlook_com(assunto, corpo_html, destinatarios, abrir_janela)
    else:
        return enviar_email_smtp_direto(assunto, corpo_html, destinatarios)


def template_email_desvio(dados: dict):
    cor_destaque = "#DC2626"
    
    html = f"""
    <!DOCTYPE html>
    <html>
    <head>
        <meta charset="utf-8">
        <style>
            body {{ font-family: 'Segoe UI', Arial, sans-serif; background-color: #f8fafc; margin: 0; padding: 20px; color: #1e293b; }}
            .container {{ max-width: 640px; background: #ffffff; margin: 0 auto; border-radius: 8px; border: 1px solid #e2e8f0; overflow: hidden; }}
            .header {{ background-color: {cor_destaque}; color: #ffffff; padding: 18px 24px; text-align: center; }}
            .content {{ padding: 24px; }}
            table {{ width: 100%; border-collapse: collapse; margin-top: 15px; }}
            td {{ padding: 10px; border-bottom: 1px solid #f1f5f9; font-size: 14px; }}
            td.campo {{ font-weight: bold; color: #64748b; width: 40%; }}
            .box-alerta {{ background-color: #fef2f2; border-left: 4px solid #ef4444; padding: 12px 16px; margin-top: 15px; border-radius: 4px; color: #991b1b; }}
            .box-justificativa {{ background-color: #fffbeb; border-left: 4px solid #f59e0b; padding: 12px 16px; margin-top: 15px; border-radius: 4px; color: #92400e; }}
        </style>
    </head>
    <body>
        <div class="container">
            <div class="header">
                <h2 style="margin: 0; font-size: 19px;">⚠️ ALERTA DE DESVIO: DUPLA BIPAGEM NA PORTARIA</h2>
                <p style="margin: 5px 0 0 0; font-size: 13px;">Nota Fiscal: {dados.get('numero_nf')} | Série: {dados.get('serie')}</p>
            </div>
            <div class="content">
                <p>Identificado um registro de <strong>dupla bipagem</strong> para a seguinte Nota Fiscal:</p>
                <table>
                    <tr>
                        <td class="campo">Número NF:</td>
                        <td><strong>{dados.get('numero_nf')}</strong> (Série: {dados.get('serie')})</td>
                    </tr>
                    <tr>
                        <td class="campo">Chave de Acesso:</td>
                        <td style="word-break: break-all; font-family: monospace; font-size: 12px;">{dados.get('chave')}</td>
                    </tr>
                    <tr>
                        <td class="campo">1ª Bipagem (Primeiro Registro):</td>
                        <td>{dados.get('data_bip1')} às {dados.get('hora_bip1')}</td>
                    </tr>
                    <tr>
                        <td class="campo" style="color: #b91c1c;">2ª Bipagem (Tentativa Duplicada):</td>
                        <td style="color: #b91c1c; font-weight: bold;">{dados.get('data_bip2')} às {dados.get('hora_bip2')}</td>
                    </tr>
                    <tr>
                        <td class="campo">Diferença de Tempo:</td>
                        <td style="font-weight: bold; font-size: 15px; color: #b91c1c;">⏱️ {dados.get('tempo_decorrido')}</td>
                    </tr>
                </table>
                <div class="box-alerta">
                    <strong>Atenção:</strong> Esta nota fiscal já possuía registro de saída anterior. A NF não deve dar saída mais de uma vez.
                </div>
    """
    if dados.get("justificativa"):
        html += f"""
                <div class="box-justificativa">
                    <strong>Justificativa Informada pela Portaria:</strong><br>
                    <em>"{dados.get('justificativa')}"</em>
                </div>
        """
    html += f"""
            </div>
            <div style="background-color: #f8fafc; padding: 12px; text-align: center; font-size: 11px; color: #94a3b8; border-top: 1px solid #e2e8f0;">
                Sistema de Portaria • Registro em {datetime.now().strftime('%d/%m/%Y %H:%M:%S')}
            </div>
        </div>
    </body>
    </html>
    """
    return html


def template_email_resumo_22h(data_ref: str, lista_notas: list):
    total = len(lista_notas)
    desvios = [n for n in lista_notas if "DESVIO" in str(n["status"]).upper()]
    regulares = [n for n in lista_notas if n["status"] == "REGULAR"]
    
    linhas = ""
    for n in lista_notas:
        chave = n["chave"]
        num_nf = n["numero_nf"]
        d1 = n["data_bip1"]
        h1 = n["hora_bip1"]
        d2 = n["data_bip2"]
        h2 = n["hora_bip2"]
        tempo = n["tempo_decorrido"]
        status = n["status"]
        just = n["justificativa"]

        is_desv = "DESVIO" in str(status).upper()
        cor_st = "#DC2626" if is_desv else "#107C41"
        txt_st = "DESVIO (DUPLO)" if is_desv else "REGULAR"
        bip2_str = f"{d2} {h2}" if d2 and h2 else "-"
        tempo_str = tempo if tempo else "-"
        just_str = just if just else "-"
        
        linhas += f"""
        <tr style="border-bottom: 1px solid #f1f5f9; text-align: center; font-size: 12px;">
            <td style="padding: 8px; font-weight: bold;">{num_nf}</td>
            <td style="padding: 8px; font-family: monospace; font-size: 11px;">{chave[:10]}...{chave[-6:]}</td>
            <td style="padding: 8px;">{d1} {h1}</td>
            <td style="padding: 8px;">{bip2_str}</td>
            <td style="padding: 8px; font-weight: bold;">{tempo_str}</td>
            <td style="padding: 8px;"><span style="background-color: {cor_st}; color: white; padding: 3px 8px; border-radius: 4px; font-size: 11px; font-weight: bold;">{txt_st}</span></td>
            <td style="padding: 8px; text-align: left; max-width: 220px; word-wrap: break-word;">{just_str}</td>
        </tr>
        """

    html = f"""
    <!DOCTYPE html>
    <html>
    <head>
        <meta charset="utf-8">
        <style>
            body {{ font-family: 'Segoe UI', Arial, sans-serif; background-color: #f8fafc; margin: 0; padding: 20px; }}
            .container {{ max-width: 900px; background: #ffffff; margin: 0 auto; border-radius: 8px; border: 1px solid #e2e8f0; overflow: hidden; }}
            .header {{ background-color: #0078D4; color: white; padding: 20px; text-align: center; }}
            .cards {{ display: flex; justify-content: space-around; background: #f8fafc; padding: 15px; border-bottom: 1px solid #e2e8f0; }}
            .card-kpi {{ text-align: center; }}
            .kpi-num {{ font-size: 22px; font-weight: bold; color: #0078D4; }}
            .kpi-leg {{ font-size: 12px; color: #64748b; }}
            table {{ width: 100%; border-collapse: collapse; }}
            th {{ background: #f1f5f9; padding: 10px; font-size: 11px; text-transform: uppercase; color: #475569; }}
        </style>
    </head>
    <body>
        <div class="container">
            <div class="header">
                <h2 style="margin: 0;">Relatório Diário de Portaria - 22:00</h2>
                <p style="margin: 5px 0 0 0; font-size: 14px;">Data de Referência: {data_ref}</p>
            </div>
            <div class="cards">
                <div class="card-kpi">
                    <div class="kpi-num">{total}</div>
                    <div class="kpi-leg">Total de NFs Movimentadas</div>
                </div>
                <div class="card-kpi">
                    <div class="kpi-num" style="color: #107C41;">{len(regulares)}</div>
                    <div class="kpi-leg">Saídas Regulares</div>
                </div>
                <div class="card-kpi">
                    <div class="kpi-num" style="color: #DC2626;">{len(desvios)}</div>
                    <div class="kpi-leg">Desvios / Duplas Bipagens</div>
                </div>
            </div>
            <div style="padding: 15px;">
                <table>
                    <thead>
                        <tr>
                            <th>Nº NF</th>
                            <th>Chave</th>
                            <th>1ª Bipagem</th>
                            <th>2ª Bipagem</th>
                            <th>Diferença</th>
                            <th>Status</th>
                            <th>Justificativa</th>
                        </tr>
                    </thead>
                    <tbody>
                        {linhas if total > 0 else '<tr><td colspan="7" style="padding: 20px; text-align: center; color: #94a3b8;">Nenhuma nota fiscal registrada hoje.</td></tr>'}
                    </tbody>
                </table>
            </div>
        </div>
    </body>
    </html>
    """
    return html


# =============================================================================
# EXPORTAÇÃO EXCEL (.XLSX)
# =============================================================================

def gerar_planilha_excel(caminho_saida: str):
    if not openpyxl:
        raise Exception("Biblioteca openpyxl não instalada. Execute 'pip install openpyxl'.")
        
    conn = obter_conexao()
    cursor = conn.cursor()
    cursor.execute("""
        SELECT chave, numero_nf, serie, data_bip1, hora_bip1, 
               data_bip2, hora_bip2, tempo_decorrido, status, justificativa, qtd_bipagens
        FROM notas 
        ORDER BY dt_completa_bip1 DESC
    """)
    registros = cursor.fetchall()
    conn.close()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Portaria - Notas"

    cabecalhos = [
        "Chave de Acesso", "Número NF", "Série", 
        "Data 1ª Bip", "Hora 1ª Bip", 
        "Data 2ª Bip", "Hora 2ª Bip", 
        "Diferença Tempo", "Status", "Justificativa da Portaria", "Qtd Bipagens"
    ]
    ws.append(cabecalhos)

    fill_cabecalho = PatternFill(start_color="0078D4", end_color="0078D4", fill_type="solid")
    font_cabecalho = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    borda_fina = Border(
        left=Side(style='thin', color='E2E8F0'),
        right=Side(style='thin', color='E2E8F0'),
        top=Side(style='thin', color='E2E8F0'),
        bottom=Side(style='thin', color='E2E8F0')
    )

    for col_idx, col_name in enumerate(cabecalhos, 1):
        c = ws.cell(row=1, column=col_idx)
        c.fill = fill_cabecalho
        c.font = font_cabecalho
        c.alignment = Alignment(horizontal="center", vertical="center")

    for row_idx, reg in enumerate(registros, 2):
        for col_idx, col_name in enumerate([
            "chave", "numero_nf", "serie", "data_bip1", "hora_bip1", 
            "data_bip2", "hora_bip2", "tempo_decorrido", "status", "justificativa", "qtd_bipagens"
        ], 1):
            val = reg[col_name]
            c = ws.cell(row=row_idx, column=col_idx, value=val if val is not None else "")
            c.border = borda_fina
            c.alignment = Alignment(horizontal="center", vertical="center")
            if col_idx in (1, 10):
                c.alignment = Alignment(horizontal="left", vertical="center")

    for col in ws.columns:
        tamanho = max(len(str(c.value or '')) for c in col)
        col_letra = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letra].width = max(min(tamanho + 3, 50), 12)

    wb.save(caminho_saida)


# =============================================================================
# AGENDAMENTO 22:00
# =============================================================================

def rotina_resumo_22h():
    hoje_str = datetime.now().strftime("%d/%m/%Y")
    conn = obter_conexao()
    cursor = conn.cursor()
    cursor.execute("""
        SELECT *
        FROM notas 
        WHERE data_bip1 = ? OR data_bip2 = ?
        ORDER BY dt_completa_bip1 ASC
    """, (hoje_str, hoje_str))
    notas_hoje = cursor.fetchall()
    conn.close()

    cfg = carregar_config()
    destinatarios = cfg.get("destinatarios", [])
    if destinatarios:
        assunto = f"📋 Relatório Consolidado de Portaria - {hoje_str} (22:00)"
        html = template_email_resumo_22h(hoje_str, notas_hoje)
        disparar_email_geral(assunto, html, destinatarios)


def worker_agendador():
    if schedule:
        schedule.every().day.at("22:00").do(rotina_resumo_22h)
        while True:
            schedule.run_pending()
            time.sleep(30)


# =============================================================================
# JANELA DE JUSTIFICATIVA DA PORTARIA (MODAL)
# =============================================================================

class ModalJustificativa(tk.Toplevel):
    def __init__(self, parent, chave, num_nf, tempo_decorrido, callback_salvar):
        super().__init__(parent)
        self.title("⚠️ DESVIO NA PORTARIA - Justificativa Obrigatória")
        self.geometry("560x380")
        self.resizable(False, False)
        self.grab_set()
        
        self.chave = chave
        self.num_nf = num_nf
        self.callback_salvar = callback_salvar

        frame_topo = tk.Frame(self, bg="#DC2626", height=50)
        frame_topo.pack(fill="x", side="top")
        
        tk.Label(
            frame_topo, 
            text="⚠️ DUPLO REGISTRO IDENTIFICADO", 
            font=("Segoe UI", 12, "bold"), 
            bg="#DC2626", 
            fg="white"
        ).pack(pady=12)

        f_corpo = tk.Frame(self, padx=20, pady=15)
        f_corpo.pack(fill="both", expand=True)

        lbl_info = tk.Label(
            f_corpo, 
            text=f"A Nota Fiscal Nº {num_nf} já registrou saída na portaria anteriormente!\n"
                 f"Diferença entre bipagens: {tempo_decorrido}\n\n"
                 f"Informe o motivo da segunda bipagem para registrar a justificativa:",
            font=("Segoe UI", 10), 
            justify="left", 
            fg="#1F2937"
        )
        lbl_info.pack(anchor="w", pady=(0, 10))

        self.txt_justificativa = tk.Text(f_corpo, height=5, font=("Segoe UI", 10), wrap="word", relief="groove", bd=2)
        self.txt_justificativa.pack(fill="both", expand=True, pady=5)
        self.txt_justificativa.focus_set()

        btn_salvar = tk.Button(
            f_corpo, 
            text="Gravar Justificativa e Atualizar Histórico", 
            font=("Segoe UI", 10, "bold"), 
            bg="#DC2626", 
            fg="white", 
            relief="flat", 
            cursor="hand2", 
            command=self.confirmar_justificativa
        )
        btn_salvar.pack(fill="x", pady=(10, 0), ipady=4)

    def confirmar_justificativa(self):
        texto = self.txt_justificativa.get("1.0", tk.END).strip()
        if not texto:
            messagebox.showwarning("Aviso", "Por favor, digite o motivo da segunda bipagem.")
            return
        
        self.callback_salvar(self.chave, texto)
        self.destroy()


# =============================================================================
# INTERFACE GRÁFICA PRINCIPAL
# =============================================================================

class AppPortariaBipagem(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("Sistema de Portaria • Controle de Saídas e Desvios de NF-e")
        self.geometry("1040x690")
        self.minsize(920, 600)
        
        self.config_data = carregar_config()
        inicializar_banco()
        
        # Variáveis de controle de Anti-Rebote
        self.em_processamento = False
        self.ultima_chave_bipada = ""
        self.ultimo_tempo_bipagem = 0.0

        self.COR_AZUL = "#0078D4"
        self.COR_VERDE = "#107C41"
        self.COR_VERMELHO = "#DC2626"
        self.COR_FUNDO = "#F3F4F6"
        
        self.configure(bg=self.COR_FUNDO)
        self.style = ttk.Style(self)
        try:
            self.style.theme_use("clam")
        except Exception:
            pass
            
        self.style.configure(".", background=self.COR_FUNDO, font=("Segoe UI", 10))
        self.style.configure("TNotebook.Tab", font=("Segoe UI", 10, "bold"), padding=[12, 6])
        self.style.configure("Treeview.Heading", font=("Segoe UI", 9, "bold"), background="#E2E8F0")
        self.style.configure("Treeview", rowheight=25, font=("Segoe UI", 9))

        self.construir_interface()
        self.carregar_dados_tabela()
        
        self.after(250, self.focar_barra_bipagem)
        
        if schedule:
            t = threading.Thread(target=worker_agendador, daemon=True)
            t.start()

    def focar_barra_bipagem(self):
        try:
            if hasattr(self, "txt_chave"):
                self.txt_chave.focus_force()
                self.txt_chave.select_range(0, tk.END)
                self.txt_chave.icursor(tk.END)
        except Exception:
            pass

    def construir_interface(self):
        top_bar = tk.Frame(self, bg=self.COR_AZUL, height=55)
        top_bar.pack(fill="x", side="top")
        
        lbl_titulo = tk.Label(
            top_bar, 
            text="🚪 PORTARIA: CONTROLE DE SAÍDA E AUDITORIA DE NF-e", 
            font=("Segoe UI", 13, "bold"), 
            bg=self.COR_AZUL, 
            fg="white"
        )
        lbl_titulo.pack(side="left", padx=20, pady=12)
        
        self.lbl_clock = tk.Label(top_bar, text="", font=("Segoe UI", 9), bg=self.COR_AZUL, fg="#E0F2FE")
        self.lbl_clock.pack(side="right", padx=20, pady=12)
        self.atualizar_relogio()

        self.notebook = ttk.Notebook(self)
        self.notebook.pack(fill="both", expand=True, padx=12, pady=10)

        self.tab_bip = ttk.Frame(self.notebook)
        self.tab_cons = ttk.Frame(self.notebook)
        self.tab_cfg = ttk.Frame(self.notebook)

        self.notebook.add(self.tab_bip, text="  📥 Bipagem na Portaria  ")
        self.notebook.add(self.tab_cons, text="  🔎 Consulta & Histórico de Desvios  ")
        self.notebook.add(self.tab_cfg, text="  ⚙️ E-mails & Configurações  ")

        self.notebook.bind("<<NotebookTabChanged>>", lambda e: self.focar_barra_bipagem())

        self.montar_aba_bipagem()
        self.montar_aba_consulta()
        self.montar_aba_config()

    def atualizar_relogio(self):
        agora = datetime.now().strftime("%d/%m/%Y %H:%M:%S")
        self.lbl_clock.config(text=f"Horário: {agora}")
        self.after(1000, self.atualizar_relogio)

    # -------------------------------------------------------------------------
    # ABA 1: BIPAGEM NA PORTARIA (SEM BOTÃO REDUNDANTE)
    # -------------------------------------------------------------------------
    def montar_aba_bipagem(self):
        card = tk.Frame(self.tab_bip, bg="white", bd=1, relief="solid")
        card.pack(fill="both", expand=True, padx=15, pady=15)

        tk.Label(
            card, 
            text="Aponte o leitor de código de barras para a Chave de Acesso (44 dígitos):", 
            font=("Segoe UI", 12, "bold"), 
            bg="white", 
            fg="#1F2937"
        ).pack(pady=(22, 10))

        f_input = tk.Frame(card, bg="white")
        f_input.pack(fill="x", padx=40, pady=8)

        self.var_chave = tk.StringVar()
        
        # Barra de Entrada Larga e com destaque
        self.txt_chave = tk.Entry(
            f_input, 
            textvariable=self.var_chave,
            font=("Consolas", 15, "bold"), 
            bd=2, 
            relief="groove", 
            justify="center", 
            highlightcolor=self.COR_AZUL, 
            highlightthickness=2
        )
        self.txt_chave.pack(side="left", fill="x", expand=True, ipady=8)

        # O leitor de código de barras dispara Enter automaticamente
        self.txt_chave.bind("<Return>", lambda e: self.processar_bipagem())
        self.txt_chave.bind("<KP_Enter>", lambda e: self.processar_bipagem())

        btn_limpar = tk.Button(
            f_input, 
            text="🔄 Limpar Barra", 
            font=("Segoe UI", 9, "bold"), 
            bg="#F3F4F6", 
            fg="#4B5563",
            relief="groove", 
            cursor="hand2", 
            command=self.limpar_barra
        )
        btn_limpar.pack(side="left", padx=(10, 0), ipady=7, ipadx=10)

        # Painel de Resposta
        self.painel_resposta = tk.Frame(card, bg="#F8FAFC", bd=1, relief="ridge")
        self.painel_resposta.pack(fill="both", expand=True, padx=40, pady=20)

        self.lbl_resp_status = tk.Label(
            self.painel_resposta, 
            text="🟢 Leitor Ativo • Aguardando Bipagem", 
            font=("Segoe UI", 13, "bold"), 
            bg="#F8FAFC", 
            fg="#107C41"
        )
        self.lbl_resp_status.pack(pady=(20, 8))

        self.lbl_resp_detalhes = tk.Label(
            self.painel_resposta, 
            text="1ª Bipagem: Registra data e hora de saída regular na portaria.\n2ª Bipagem da mesma NF: Registra desvio, calcula tempo decorrido e envia e-mail de alerta.", 
            font=("Segoe UI", 10), 
            bg="#F8FAFC", 
            fg="#475569", 
            justify="center"
        )
        self.lbl_resp_detalhes.pack(pady=5, padx=20)

    def limpar_barra(self):
        self.var_chave.set("")
        self.ultima_chave_bipada = ""
        self.ultimo_tempo_bipagem = 0.0
        self.focar_barra_bipagem()

    def processar_bipagem(self):
        if self.em_processamento:
            return
            
        chave_digitada = self.var_chave.get().strip()
        if not chave_digitada:
            self.focar_barra_bipagem()
            return

        dados = sanitizar_e_extrair_chave(chave_digitada)
        if not dados["valida"]:
            self.exibir_status_painel(
                "❌ Chave Inválida / Incompleta", 
                f"O código lido possui {len(dados['chave'])} dígitos numéricos.\n"
                f"Uma chave de NF-e precisa ter exatamente 44 dígitos.\n"
                f"Conteúdo lido: {dados['chave']}", 
                "#FEF2F2", "#B91C1C"
            )
            self.focar_barra_bipagem()
            return

        chave = dados["chave"]
        num_nf = dados["numero_nf"]
        serie = dados["serie"]
        agora_timestamp = time.time()

        # Anti-rebote (ignora comandos duplicados em menos de 2.5s)
        if chave == self.ultima_chave_bipada and (agora_timestamp - self.ultimo_tempo_bipagem < 2.5):
            self.focar_barra_bipagem()
            return

        self.em_processamento = True
        self.ultima_chave_bipada = chave
        self.ultimo_tempo_bipagem = agora_timestamp

        # Mantém a chave na barra e seleciona
        self.var_chave.set(chave)

        agora = datetime.now()
        dt_str = agora.strftime("%d/%m/%Y")
        hr_str = agora.strftime("%H:%M:%S")
        iso_str = agora.isoformat()

        try:
            conn = obter_conexao()
            c = conn.cursor()
            c.execute("SELECT * FROM notas WHERE chave = ?", (chave,))
            nota = c.fetchone()

            if not nota:
                # =============================================================
                # 1ª BIPAGEM - SAÍDA REGULAR (SEM E-MAIL DE DESVIO)
                # =============================================================
                c.execute("""
                    INSERT INTO notas (
                        chave, numero_nf, serie, data_bip1, hora_bip1, 
                        dt_completa_bip1, status, justificativa, qtd_bipagens
                    ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, 1)
                """, (chave, num_nf, serie, dt_str, hr_str, iso_str, "REGULAR", ""))
                conn.commit()
                conn.close()

                self.exibir_status_painel(
                    "✅ 1ª Bipagem: Saída Liberada e Gravada",
                    f"📄 Nota Fiscal: {num_nf} (Série {serie})\n"
                    f"🔑 Chave: {chave}\n"
                    f"📅 Data/Hora: {dt_str} às {hr_str}\n\n"
                    f"Status: Salvo no banco de dados. Registro regular de portaria.",
                    "#EFF6FF", "#1D4ED8"
                )

            else:
                # =============================================================
                # 2ª BIPAGEM - DESVIO DETECTADO (DISPARA E-MAIL E MODAL)
                # =============================================================
                dt1_iso = nota["dt_completa_bip1"] or iso_str
                try:
                    dt1 = datetime.fromisoformat(dt1_iso)
                except Exception:
                    dt1 = agora

                tempo_txt, mins = calcular_diferenca_tempo(dt1, agora)
                
                try:
                    qtd_bip = int(nota["qtd_bipagens"])
                except Exception:
                    qtd_bip = 1
                qtd_atual = qtd_bip + 1

                c.execute("""
                    UPDATE notas SET 
                        data_bip2 = ?, hora_bip2 = ?, dt_completa_bip2 = ?,
                        tempo_decorrido = ?, minutos_decorridos = ?,
                        status = ?, qtd_bipagens = ?
                    WHERE chave = ?
                """, (dt_str, hr_str, iso_str, tempo_txt, mins, "DESVIO REGISTRADO", qtd_atual, chave))
                conn.commit()
                conn.close()

                payload = {
                    "chave": chave, "numero_nf": num_nf, "serie": serie,
                    "data_bip1": nota["data_bip1"], "hora_bip1": nota["hora_bip1"],
                    "data_bip2": dt_str, "hora_bip2": hr_str,
                    "tempo_decorrido": tempo_txt,
                    "justificativa": nota["justificativa"] or "Aguardando preenchimento da portaria"
                }

                self.exibir_status_painel(
                    "🚨 ALERTA: DESVIO DETECTADO (DUPLA BIPAGEM NA PORTARIA)",
                    f"⚠️ ATENÇÃO: A NF {num_nf} já registrou saída anteriormente!\n\n"
                    f"⏱️ 1ª Bipagem: {nota['data_bip1']} às {nota['hora_bip1']}\n"
                    f"⏱️ 2ª Bipagem (Tentativa Duplicada): {dt_str} às {hr_str}\n"
                    f"⏳ TEMPO DECORRIDO ENTRE AS BIPAGENS: {tempo_txt}\n\n"
                    f"E-mail de desvio enviado à gestão. Preencha a justificativa na janela.",
                    "#FEF2F2", "#B91C1C"
                )

                # Dispara o e-mail de alerta de desvio
                threading.Thread(target=self.disparar_email_desvio_bg, args=(payload,), daemon=True).start()

                # Abre a tela modal de justificativa
                ModalJustificativa(self, chave, num_nf, tempo_txt, self.salvar_justificativa_portaria)

        except Exception as e:
            messagebox.showerror("Erro de Banco de Dados", f"Falha ao salvar no SQLite:\n{str(e)}")
        finally:
            self.em_processamento = False
            self.carregar_dados_tabela()
            self.focar_barra_bipagem()

    def salvar_justificativa_portaria(self, chave, texto_justificativa):
        try:
            conn = obter_conexao()
            c = conn.cursor()
            c.execute("UPDATE notas SET justificativa = ? WHERE chave = ?", (texto_justificativa, chave))
            c.execute("SELECT * FROM notas WHERE chave = ?", (chave,))
            row = c.fetchone()
            conn.commit()
            conn.close()

            if row:
                payload = {
                    "chave": chave, "numero_nf": row["numero_nf"], "serie": row["serie"],
                    "data_bip1": row["data_bip1"], "hora_bip1": row["hora_bip1"],
                    "data_bip2": row["data_bip2"], "hora_bip2": row["hora_bip2"],
                    "tempo_decorrido": row["tempo_decorrido"],
                    "justificativa": texto_justificativa
                }
                threading.Thread(target=self.disparar_email_desvio_bg, args=(payload, True), daemon=True).start()

            messagebox.showinfo("Justificativa Salva", "A justificativa da portaria foi gravada no banco de dados e enviada por e-mail com sucesso!")
        except Exception as e:
            messagebox.showerror("Erro ao Salvar Justificativa", str(e))
        finally:
            self.carregar_dados_tabela()
            self.focar_barra_bipagem()

    def disparar_email_desvio_bg(self, dados: dict, com_justificativa=False):
        cfg = carregar_config()
        destinatarios = cfg.get("destinatarios", [])
        if destinatarios:
            sufixo = " (Com Justificativa)" if com_justificativa else ""
            assunto = f"[ALERTA DE DESVIO{sufixo}] NF {dados.get('numero_nf')} - {datetime.now().strftime('%d/%m/%Y %H:%M')}"
            corpo = template_email_desvio(dados)
            sucesso, msg = disparar_email_geral(assunto, corpo, destinatarios)
            if not sucesso:
                print(f"[ALERTA DE ENVIO] Falha no disparo de e-mail: {msg}")

    def exibir_status_painel(self, titulo, detalhes, bg_cor, fg_cor):
        self.painel_resposta.config(bg=bg_cor)
        self.lbl_resp_status.config(text=titulo, bg=bg_cor, fg=fg_cor)
        self.lbl_resp_detalhes.config(text=detalhes, bg=bg_cor)

    # -------------------------------------------------------------------------
    # ABA 2: CONSULTA & HISTÓRICO DE DESVIOS
    # -------------------------------------------------------------------------
    def montar_aba_consulta(self):
        card = tk.Frame(self.tab_cons, bg="white", bd=1, relief="solid")
        card.pack(fill="both", expand=True, padx=15, pady=15)

        f_busca = tk.Frame(card, bg="white")
        f_busca.pack(fill="x", padx=15, pady=10)

        tk.Label(f_busca, text="Buscar Chave / NF:", font=("Segoe UI", 9, "bold"), bg="white").pack(side="left", padx=(0, 5))
        self.txt_busca = tk.Entry(f_busca, font=("Segoe UI", 9), width=28)
        self.txt_busca.pack(side="left", padx=5)
        self.txt_busca.bind("<KeyRelease>", lambda e: self.carregar_dados_tabela())

        btn_atualizar = tk.Button(f_busca, text="🔄 Atualizar", bg="#E2E8F0", command=self.carregar_dados_tabela)
        btn_atualizar.pack(side="left", padx=10)

        btn_excel = tk.Button(
            f_busca, 
            text="📊 Exportar para Excel (.xlsx)", 
            bg=self.COR_VERDE, 
            fg="white", 
            font=("Segoe UI", 9, "bold"), 
            relief="flat", 
            cursor="hand2",
            command=self.acao_exportar
        )
        btn_excel.pack(side="right")

        colunas = ("chave", "nf", "serie", "d1", "h1", "d2", "h2", "tempo", "status", "justificativa")
        self.grid = ttk.Treeview(card, columns=colunas, show="headings", selectmode="browse")

        self.grid.heading("chave", text="Chave de Acesso")
        self.grid.heading("nf", text="Nº NF")
        self.grid.heading("serie", text="Série")
        self.grid.heading("d1", text="Data 1ª Bip")
        self.grid.heading("h1", text="Hora 1ª Bip")
        self.grid.heading("d2", text="Data 2ª Bip")
        self.grid.heading("h2", text="Hora 2ª Bip")
        self.grid.heading("tempo", text="Diferença")
        self.grid.heading("status", text="Status")
        self.grid.heading("justificativa", text="Justificativa da Portaria")

        self.grid.column("chave", width=220, anchor="w")
        self.grid.column("nf", width=70, anchor="center")
        self.grid.column("serie", width=45, anchor="center")
        self.grid.column("d1", width=80, anchor="center")
        self.grid.column("h1", width=70, anchor="center")
        self.grid.column("d2", width=80, anchor="center")
        self.grid.column("h2", width=70, anchor="center")
        self.grid.column("tempo", width=95, anchor="center")
        self.grid.column("status", width=110, anchor="center")
        self.grid.column("justificativa", width=220, anchor="w")

        scroll = ttk.Scrollbar(card, orient="vertical", command=self.grid.yview)
        self.grid.configure(yscrollcommand=scroll.set)

        self.grid.pack(side="left", fill="both", expand=True, padx=(15, 0), pady=10)
        scroll.pack(side="right", fill="y", padx=(0, 15), pady=10)

    def carregar_dados_tabela(self):
        filtro = self.txt_busca.get().strip() if hasattr(self, "txt_busca") else ""
        for i in self.grid.get_children():
            self.grid.delete(i)

        try:
            conn = obter_conexao()
            c = conn.cursor()
            if filtro:
                c.execute("""
                    SELECT *
                    FROM notas 
                    WHERE chave LIKE ? OR numero_nf LIKE ?
                    ORDER BY dt_completa_bip1 DESC
                """, (f"%{filtro}%", f"%{filtro}%"))
            else:
                c.execute("""
                    SELECT *
                    FROM notas 
                    ORDER BY dt_completa_bip1 DESC
                """)

            for row in c.fetchall():
                st_txt = "🚨 DESVIO" if "DESVIO" in str(row["status"]).upper() else "REGULAR"
                self.grid.insert("", "end", values=(
                    row["chave"], row["numero_nf"], row["serie"], 
                    row["data_bip1"], row["hora_bip1"], 
                    row["data_bip2"] or "-", row["hora_bip2"] or "-", 
                    row["tempo_decorrido"] or "-", st_txt,
                    row["justificativa"] or "-"
                ))
            conn.close()
        except Exception as e:
            print(f"Erro ao carregar tabela: {e}")

    def acao_exportar(self):
        caminho = filedialog.asksaveasfilename(
            defaultextension=".xlsx",
            filetypes=[("Arquivos Excel", "*.xlsx")],
            initialfile=f"Relatorio_Portaria_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        )
        if caminho:
            try:
                gerar_planilha_excel(caminho)
                messagebox.showinfo("Sucesso", f"Planilha salva com sucesso em:\n{caminho}")
            except Exception as e:
                messagebox.showerror("Erro", f"Falha ao gerar Excel: {e}")

    # -------------------------------------------------------------------------
    # ABA 3: CONFIGURAÇÕES
    # -------------------------------------------------------------------------
    def montar_aba_config(self):
        card = tk.Frame(self.tab_cfg, bg="white", bd=1, relief="solid")
        card.pack(fill="both", expand=True, padx=15, pady=15)

        # Lado Esquerdo: Configurações de Envio
        f_smtp = tk.LabelFrame(card, text=" Configuração de Envio de E-mail ", font=("Segoe UI", 9, "bold"), bg="white", padx=12, pady=12)
        f_smtp.place(relx=0.02, rely=0.03, relwidth=0.46, relheight=0.94)

        tk.Label(f_smtp, text="Método de Envio:", font=("Segoe UI", 9, "bold"), bg="white").pack(anchor="w", pady=(0, 2))
        
        self.var_metodo = tk.StringVar(value=self.config_data.get("metodo_envio", "OUTLOOK_APP"))
        
        rb_outlook = tk.Radiobutton(
            f_smtp, 
            text="📧 Outlook Desktop (MAPI / PowerShell Integrado)", 
            variable=self.var_metodo, 
            value="OUTLOOK_APP", 
            bg="white",
            font=("Segoe UI", 9)
        )
        rb_outlook.pack(anchor="w", pady=2)

        self.var_abrir_janela = tk.BooleanVar(value=self.config_data.get("exibir_janela_email", False))
        chk_janela = tk.Checkbutton(
            f_smtp,
            text="👁️ Abrir janela do e-mail na tela antes de enviar (Para Teste)",
            variable=self.var_abrir_janela,
            bg="white",
            font=("Segoe UI", 8, "italic")
        )
        chk_janela.pack(anchor="w", padx=20, pady=(0, 8))

        rb_smtp = tk.Radiobutton(
            f_smtp, 
            text="🌐 SMTP Office 365 (Requer SMTP liberado pelo TI)", 
            variable=self.var_metodo, 
            value="SMTP", 
            bg="white",
            font=("Segoe UI", 9)
        )
        rb_smtp.pack(anchor="w", pady=(0, 8))

        tk.Label(f_smtp, text="E-mail Remetente (Opcional p/ Outlook Desktop):", bg="white").pack(anchor="w", pady=(2, 2))
        self.txt_remetente = tk.Entry(f_smtp, width=32)
        self.txt_remetente.insert(0, self.config_data.get("email_remetente", ""))
        self.txt_remetente.pack(fill="x", pady=(0, 6))

        tk.Label(f_smtp, text="Senha de Aplicativo (Apenas se usar SMTP):", bg="white").pack(anchor="w", pady=(2, 2))
        self.txt_senha = tk.Entry(f_smtp, show="•", width=32)
        self.txt_senha.insert(0, self.config_data.get("senha_app", ""))
        self.txt_senha.pack(fill="x", pady=(0, 6))

        tk.Label(f_smtp, text="Servidor SMTP:", bg="white").pack(anchor="w", pady=(2, 2))
        self.txt_servidor = tk.Entry(f_smtp, width=32)
        self.txt_servidor.insert(0, self.config_data.get("smtp_server", "smtp.office365.com"))
        self.txt_servidor.pack(fill="x", pady=(0, 6))

        tk.Label(f_smtp, text="Porta SMTP:", bg="white").pack(anchor="w", pady=(2, 2))
        self.txt_porta = tk.Entry(f_smtp, width=8)
        self.txt_porta.insert(0, str(self.config_data.get("smtp_port", 587)))
        self.txt_porta.pack(anchor="w", pady=(0, 10))

        btn_testar = tk.Button(f_smtp, text="✉️ Testar Envio de E-mail Agora", bg="#F1F5F9", font=("Segoe UI", 9, "bold"), command=self.testar_envio)
        btn_testar.pack(fill="x", pady=4)

        btn_resumo = tk.Button(f_smtp, text="🚀 Disparar Resumo das 22h Agora", bg="#F1F5F9", command=rotina_resumo_22h)
        btn_resumo.pack(fill="x", pady=4)

        # Lado Direito: Destinatários
        f_dest = tk.LabelFrame(card, text=" Grupo de E-mails Destinatários (Auditoria/Gestão) ", font=("Segoe UI", 9, "bold"), bg="white", padx=12, pady=12)
        f_dest.place(relx=0.51, rely=0.03, relwidth=0.46, relheight=0.94)

        f_add = tk.Frame(f_dest, bg="white")
        f_add.pack(fill="x", pady=(0, 8))

        self.txt_novo_email = tk.Entry(f_add, font=("Segoe UI", 9))
        self.txt_novo_email.pack(side="left", fill="x", expand=True, padx=(0, 5))

        btn_add = tk.Button(f_add, text="➕ Adicionar", bg=self.COR_AZUL, fg="white", font=("Segoe UI", 8, "bold"), command=self.adicionar_email)
        btn_add.pack(side="left")

        self.lista_emails = tk.Listbox(f_dest, height=9, font=("Segoe UI", 9))
        self.lista_emails.pack(fill="both", expand=True, pady=(0, 8))
        for email in self.config_data.get("destinatarios", []):
            self.lista_emails.insert(tk.END, email)

        btn_del = tk.Button(f_dest, text="🗑️ Remover E-mail Selecionado", bg="#EF4444", fg="white", command=self.remover_email)
        btn_del.pack(fill="x", pady=(0, 8))

        btn_salvar = tk.Button(f_dest, text="💾 Salvar Todas as Configurações", bg=self.COR_VERDE, fg="white", font=("Segoe UI", 9, "bold"), command=self.salvar_configs)
        btn_salvar.pack(fill="x", ipady=3)

    def adicionar_email(self):
        email = self.txt_novo_email.get().strip()
        if email and "@" in email and "." in email:
            self.lista_emails.insert(tk.END, email)
            self.txt_novo_email.delete(0, tk.END)
        else:
            messagebox.showwarning("Aviso", "Informe um endereço de e-mail válido.")

    def remover_email(self):
        sel = self.lista_emails.curselection()
        if sel:
            self.lista_emails.delete(sel[0])

    def salvar_configs(self):
        nova_config = {
            "metodo_envio": self.var_metodo.get(),
            "exibir_janela_email": self.var_abrir_janela.get(),
            "smtp_server": self.txt_servidor.get().strip(),
            "smtp_port": int(self.txt_porta.get().strip() or 587),
            "email_remetente": self.txt_remetente.get().strip(),
            "senha_app": self.txt_senha.get().strip(),
            "destinatarios": list(self.lista_emails.get(0, tk.END)),
            "horario_resumo": "22:00"
        }
        salvar_config(nova_config)
        self.config_data = nova_config
        messagebox.showinfo("Sucesso", "Configurações e lista de destinatários salvas com sucesso!")

    def testar_envio(self):
        self.salvar_configs()
        destinatarios = self.config_data.get("destinatarios", [])
        if not destinatarios:
            messagebox.showwarning("Aviso", "Adicione ao menos um e-mail na lista de destinatários à direita.")
            return

        ok, msg = disparar_email_geral(
            assunto="🧪 Teste de Conexão - Sistema de Portaria de NF-e",
            corpo_html="<h3>Teste realizado com sucesso!</h3><p>O sistema de portaria está configurado e pronto para envio.</p>",
            destinatarios=destinatarios
        )
        if ok:
            messagebox.showinfo("Diagnóstico de Envio", f"{msg}\n\nDestinatários: {', '.join(destinatarios)}")
        else:
            messagebox.showerror("Erro de Envio", msg)


# =============================================================================
# INICIALIZAÇÃO
# =============================================================================

if __name__ == "__main__":
    try:
        app = AppPortariaBipagem()
        app.mainloop()
    except Exception as e:
        import traceback
        print("\n[ERRO AO INICIAR APLICAÇÃO]:")
        traceback.print_exc()
        input("\nPressione Enter para sair...")
