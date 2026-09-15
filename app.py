# -*- coding: utf-8 -*-
import os
import re
import json
import sqlite3
import smtplib
import threading
from io import BytesIO
from datetime import datetime, date, time, timezone, timedelta

import openpyxl
from flask import Flask, jsonify, render_template_string, request, send_file
from email.mime.text import MIMEText
from email.mime.multipart import MIMEMultipart

# =========================================================
# APP / FUSO
# =========================================================
try:
    from zoneinfo import ZoneInfo
    LOCAL_TZ = ZoneInfo(os.getenv("APP_TZ", "America/Sao_Paulo"))
except Exception:
    LOCAL_TZ = timezone(timedelta(hours=-3))

app = Flask(__name__)
application = app  # gunicorn app:application

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DB_NAME = os.path.join(BASE_DIR, "bipagem_nfe.db")
CONFIG_FILE = os.path.join(BASE_DIR, "config_bipagem.json")


# =========================================================
# HELPERS
# =========================================================
def agora_local() -> datetime:
    return datetime.now(LOCAL_TZ)


def only_digits(s) -> str:
    return re.sub(r"\D", "", str(s or ""))


def normalizar_header(s: str) -> str:
    s = str(s or "").strip().lower()
    s = re.sub(r"\s+", " ", s)
    return s


def parse_iso_db(value: str):
    if not value:
        return None
    try:
        dt = datetime.fromisoformat(value)
    except Exception:
        return None

    if dt.tzinfo is None:
        dt = dt.replace(tzinfo=LOCAL_TZ)
    else:
        dt = dt.astimezone(LOCAL_TZ)
    return dt


def parse_data_hora_ref(data_ref, hora_ref):
    """
    Converte data/hora vinda da base de comparação (Excel/Texto)
    para datetime com fuso LOCAL_TZ.
    """
    if data_ref in (None, ""):
        return None

    # Caso Excel tenha vindo como datetime/date
    if isinstance(data_ref, datetime):
        d = data_ref.date()
    elif isinstance(data_ref, date):
        d = data_ref
    else:
        d = None

    # Hora pode vir como datetime/time/str
    if isinstance(hora_ref, datetime):
        h = hora_ref.time().replace(microsecond=0)
    elif isinstance(hora_ref, time):
        h = hora_ref.replace(microsecond=0)
    elif isinstance(hora_ref, str) and hora_ref.strip():
        txt = hora_ref.strip()
        for fmt in ("%H:%M:%S", "%H:%M"):
            try:
                h = datetime.strptime(txt, fmt).time()
                break
            except Exception:
                h = None
    else:
        h = None

    if d is not None:
        if h is None:
            h = time(0, 0, 0)
        return datetime.combine(d, h).replace(tzinfo=LOCAL_TZ)

    # Se data veio como string
    data_txt = str(data_ref).strip()
    hora_txt = str(hora_ref or "").strip() or "00:00:00"

    candidatos = [
        f"{data_txt} {hora_txt}",
        f"{data_txt} 00:00:00",
        data_txt,  # pode vir só data
    ]
    formatos = [
        "%d/%m/%Y %H:%M:%S",
        "%d/%m/%Y %H:%M",
        "%Y-%m-%d %H:%M:%S",
        "%Y-%m-%d %H:%M",
        "%d/%m/%Y",
        "%Y-%m-%d",
    ]

    for txt in candidatos:
        for fmt in formatos:
            try:
                dt = datetime.strptime(txt, fmt)
                return dt.replace(tzinfo=LOCAL_TZ)
            except Exception:
                pass
    return None


def calcular_diferenca(dt1: datetime, dt2: datetime):
    delta = dt2 - dt1
    s = max(0, int(delta.total_seconds()))
    d, r = divmod(s, 86400)
    h, r = divmod(r, 3600)
    m, seg = divmod(r, 60)

    partes = []
    if d > 0:
        partes.append(f"{d}d")
    if h > 0 or d > 0:
        partes.append(f"{h}h")
    partes.append(f"{m}m {seg}s")
    return " ".join(partes), round(s / 60.0, 2)


def sanitizar_e_extrair_chave(texto):
    nums = only_digits(texto)
    if len(nums) >= 44:
        chave = nums[:44]
        try:
            serie = str(int(chave[22:25]))
        except Exception:
            serie = chave[22:25]
        try:
            numero_nf = str(int(chave[25:34]))
        except Exception:
            numero_nf = chave[25:34]
        return {"valida": True, "chave": chave, "numero_nf": numero_nf, "serie": serie}
    return {"valida": False, "chave": nums, "numero_nf": "N/D", "serie": "N/D"}


# =========================================================
# BANCO
# =========================================================
def obter_conexao():
    conn = sqlite3.connect(DB_NAME)
    conn.row_factory = sqlite3.Row
    return conn


def inicializar_banco():
    conn = obter_conexao()
    c = conn.cursor()

    # Tabela operacional (visível no histórico)
    c.execute("""
        CREATE TABLE IF NOT EXISTS notas (
            chave TEXT PRIMARY KEY,
            numero_nf TEXT,
            serie TEXT,
            data_bip1 TEXT,
            hora_bip1 TEXT,
            dt_completa_bip1 TEXT,
            data_bip2 TEXT,
            hora_bip2 TEXT,
            dt_completa_bip2 TEXT,
            tempo_decorrido TEXT,
            minutos_decorridos REAL,
            status TEXT,
            justificativa TEXT,
            qtd_bipagens INTEGER DEFAULT 1,
            observacao TEXT,
            email_desvio_enviado INTEGER DEFAULT 0
        )
    """)

    # Tabela oculta de comparação (não exibida no histórico)
    c.execute("""
        CREATE TABLE IF NOT EXISTS base_comparacao (
            chave TEXT PRIMARY KEY,
            numero_nf TEXT,
            serie TEXT,
            data_ref TEXT,
            hora_ref TEXT,
            origem_arquivo TEXT,
            importado_em TEXT
        )
    """)

    c.execute("CREATE INDEX IF NOT EXISTS idx_base_comp_nf ON base_comparacao(numero_nf)")
    c.execute("CREATE INDEX IF NOT EXISTS idx_base_comp_serie ON base_comparacao(serie)")

    # Migração de colunas antigas na tabela notas
    c.execute("PRAGMA table_info(notas)")
    cols = [x["name"] for x in c.fetchall()]

    needed = {
        "data_bip2": "TEXT",
        "hora_bip2": "TEXT",
        "dt_completa_bip2": "TEXT",
        "tempo_decorrido": "TEXT",
        "minutos_decorridos": "REAL",
        "status": "TEXT",
        "justificativa": "TEXT",
        "qtd_bipagens": "INTEGER DEFAULT 1",
        "observacao": "TEXT",
        "email_desvio_enviado": "INTEGER DEFAULT 0",
    }

    for col, typ in needed.items():
        if col not in cols:
            try:
                c.execute(f"ALTER TABLE notas ADD COLUMN {col} {typ}")
            except Exception:
                pass

    conn.commit()
    conn.close()


# =========================================================
# E-MAIL
# =========================================================
def carregar_config():
    if os.path.exists(CONFIG_FILE):
        try:
            with open(CONFIG_FILE, "r", encoding="utf-8") as f:
                return json.load(f)
        except Exception:
            pass
    return {
        "destinatarios": [],
        "smtp_server": "smtp.office365.com",
        "smtp_port": 587,
        "email_remetente": "",
        "senha_app": "",
    }


def enviar_email_smtp(assunto, corpo_html, destinatarios):
    cfg = carregar_config()
    remetente = cfg.get("email_remetente", "").strip()
    senha = cfg.get("senha_app", "").strip()
    servidor = cfg.get("smtp_server", "smtp.office365.com").strip()

    try:
        porta = int(cfg.get("smtp_port", 587))
    except Exception:
        porta = 587

    dests = [d.strip() for d in (destinatarios or []) if d and d.strip()]
    if not remetente or not senha or not dests:
        return False

    try:
        msg = MIMEMultipart("alternative")
        msg["Subject"] = assunto
        msg["From"] = remetente
        msg["To"] = ", ".join(dests)
        msg.attach(MIMEText(corpo_html, "html", "utf-8"))

        with smtplib.SMTP(servidor, porta, timeout=20) as server:
            server.starttls()
            server.login(remetente, senha)
            server.sendmail(remetente, dests, msg.as_string())
        return True
    except Exception:
        return False


def template_email_desvio(dados):
    return f"""
    <div style="font-family:Arial,sans-serif;padding:20px;color:#333">
      <div style="background:#DC2626;color:#fff;padding:15px;border-radius:6px;text-align:center">
        <h2 style="margin:0">⚠️ ALERTA DE DESVIO: DUPLA BIPAGEM NA PORTARIA</h2>
        <p style="margin:5px 0 0 0">Nota Fiscal: {dados.get('numero_nf')} | Série: {dados.get('serie')}</p>
      </div>
      <div style="padding:20px;border:1px solid #e5e7eb;border-radius:6px;margin-top:15px">
        <p><strong>Chave:</strong> <code style="font-size:11px">{dados.get('chave')}</code></p>
        <p><strong>1ª Bipagem:</strong> {dados.get('data_bip1')} às {dados.get('hora_bip1')}</p>
        <p><strong>2ª Bipagem:</strong> <span style="color:#DC2626;font-weight:bold">{dados.get('data_bip2')} às {dados.get('hora_bip2')}</span></p>
        <p><strong>Diferença:</strong> ⏱️ <b>{dados.get('tempo_decorrido')}</b></p>
        <p><strong>Justificativa:</strong> {dados.get('justificativa', 'Pendente')}</p>
      </div>
    </div>
    """


# =========================================================
# IMPORTAÇÃO BASE OCULTA
# =========================================================
def detectar_colunas_base(ws):
    headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
    hnorm = [normalizar_header(h) for h in headers]

    def find_col(cands):
        for i, h in enumerate(hnorm, start=1):
            for c in cands:
                if h == c or c in h:
                    return i
        return None

    return {
        "chave": find_col(["chave", "chave de acesso", "chave_acesso"]),
        "numero_nf": find_col(["numero nf", "número nf", "numero_nf", "nf"]),
        "serie": find_col(["serie", "série"]),
        "data": find_col(["data"]),
        "hora": find_col(["hora"]),
    }


def extrair_registros_base(ws, colmap):
    validos = []
    rejeitados = []
    vistos = set()

    for r in range(2, ws.max_row + 1):
        raw = ws.cell(r, colmap["chave"]).value if colmap["chave"] else None
        chave = only_digits(raw)

        if not chave:
            rejeitados.append({"linha": r, "motivo": "CHAVE_VAZIA", "valor": raw})
            continue

        if len(chave) != 44:
            rejeitados.append({"linha": r, "motivo": "CHAVE_INVALIDA_TAMANHO", "valor": raw, "digitos": len(chave)})
            continue

        if chave in vistos:
            # duplicidade no arquivo de carga
            continue
        vistos.add(chave)

        nf = ws.cell(r, colmap["numero_nf"]).value if colmap["numero_nf"] else None
        serie = ws.cell(r, colmap["serie"]).value if colmap["serie"] else None
        data_ref = ws.cell(r, colmap["data"]).value if colmap["data"] else None
        hora_ref = ws.cell(r, colmap["hora"]).value if colmap["hora"] else None

        # NF
        if nf is not None:
            if isinstance(nf, (int, float)):
                nf = str(int(nf))
            else:
                nf = str(nf).strip()
        else:
            nf = ""

        # Série
        if serie in (None, ""):
            serie = chave[22:25]
            try:
                serie = str(int(serie))
            except Exception:
                pass
        else:
            if isinstance(serie, (int, float)):
                serie = str(int(serie))
            else:
                serie = str(serie).strip()

        # Data/Hora em texto
        if isinstance(data_ref, datetime):
            data_txt = data_ref.strftime("%d/%m/%Y")
        elif isinstance(data_ref, date):
            data_txt = data_ref.strftime("%d/%m/%Y")
        elif data_ref is None:
            data_txt = ""
        else:
            data_txt = str(data_ref).strip()

        if isinstance(hora_ref, datetime):
            hora_txt = hora_ref.strftime("%H:%M:%S")
        elif isinstance(hora_ref, time):
            hora_txt = hora_ref.strftime("%H:%M:%S")
        elif hora_ref is None:
            hora_txt = ""
        else:
            hora_txt = str(hora_ref).strip()

        validos.append({
            "chave": chave,
            "numero_nf": nf,
            "serie": serie,
            "data_ref": data_txt,
            "hora_ref": hora_txt,
        })

    return validos, rejeitados


# =========================================================
# HTML
# =========================================================
HTML = """<!DOCTYPE html>
<html lang="pt-BR">
<head>
  <meta charset="UTF-8">
  <meta name="viewport" content="width=device-width, initial-scale=1.0">
  <title>Portaria Web • Controle de NF-e</title>
  <link href="https://cdn.jsdelivr.net/npm/bootstrap@5.3.0/dist/css/bootstrap.min.css" rel="stylesheet">
  <style>
    body{background:#f3f4f6;font-family:'Segoe UI',sans-serif;}
    .barra-input{font-size:1.4rem;font-family:monospace;font-weight:bold;}
    .top-filtro .form-control{max-width:360px;}
  </style>
</head>
<body>
  <nav class="navbar navbar-dark bg-primary mb-4 shadow-sm">
    <div class="container-fluid px-4">
      <span class="navbar-brand fw-bold">🚪 PORTARIA WEB • CONTROLE DE SAÍDAS DE NF-e</span>
      <span class="text-white-50 fw-bold" id="clock"></span>
    </div>
  </nav>

  <div class="container-fluid px-4">
    <div class="card p-4 mb-4 border-0 shadow-sm">
      <h5 class="text-secondary fw-bold mb-3">⚡ Leitura de Código de Barras (Chave de 44 dígitos)</h5>
      <div class="input-group mb-3">
        <input type="text" id="chaveInput" class="form-control barra-input text-center"
               placeholder="Aponte o leitor de código de barras aqui..." autofocus autocomplete="off">
        <button id="btnRegistrar" type="button" class="btn btn-primary px-4 fw-bold">Registrar</button>
        <button id="btnLimpar" type="button" class="btn btn-outline-secondary px-4 fw-bold">Limpar</button>
      </div>
      <div id="painel" class="alert alert-light border text-center my-0 py-3 fw-bold text-success">
        🟢 Leitor Pronto para Bipagem
      </div>
    </div>

    <div class="card p-4 border-0 shadow-sm">
      <div class="d-flex justify-content-between align-items-center mb-3">
        <h5 class="text-secondary fw-bold mb-0">🔎 Histórico de Registros e Desvios</h5>
        <a href="/api/exportar" class="btn btn-success fw-bold">📊 Baixar Excel (.xlsx)</a>
      </div>

      <div class="d-flex align-items-center gap-2 mb-3 top-filtro">
        <label for="buscaInput" class="fw-bold text-secondary mb-0">Buscar Chave / NF:</label>
        <input id="buscaInput" type="text" class="form-control" placeholder="Digite a chave ou número da NF">
        <button id="btnAtualizar" type="button" class="btn btn-outline-secondary fw-bold">Atualizar</button>
      </div>

      <div class="table-responsive">
        <table class="table table-hover align-middle">
          <thead class="table-light">
            <tr class="text-center text-secondary small">
              <th>NF</th>
              <th>Chave de Acesso</th>
              <th>Série</th>
              <th>Data 1ª Bip</th>
              <th>Hora 1ª Bip</th>
              <th>Data 2ª Bip</th>
              <th>Hora 2ª Bip</th>
              <th>Diferença</th>
              <th>Status</th>
              <th>Justificativa</th>
            </tr>
          </thead>
          <tbody id="corpo"></tbody>
        </table>
      </div>
    </div>
  </div>

  <div class="modal fade" id="modalJust" tabindex="-1" data-bs-backdrop="static">
    <div class="modal-dialog modal-dialog-centered"><div class="modal-content">
      <div class="modal-header bg-danger text-white">
        <h5 class="modal-title fw-bold">⚠️ DUPLO REGISTRO IDENTIFICADO</h5>
      </div>
      <div class="modal-body">
        <p id="descModal" class="mb-3"></p>
        <label class="form-label fw-bold">Informe o motivo da 2ª bipagem (Obrigatório):</label>
        <textarea id="justInput" class="form-control" rows="3" placeholder="Digite aqui..."></textarea>
      </div>
      <div class="modal-footer">
        <button id="btnSalvarJust" type="button" class="btn btn-danger fw-bold w-100">
          Gravar Justificativa e Enviar E-mail
        </button>
      </div>
    </div></div>
  </div>

  <script src="https://cdn.jsdelivr.net/npm/bootstrap@5.3.0/dist/js/bootstrap.bundle.min.js"></script>
  <script>
    document.addEventListener('DOMContentLoaded', function () {
      const inputElem = document.getElementById('chaveInput');
      const btnRegistrar = document.getElementById('btnRegistrar');
      const btnLimpar = document.getElementById('btnLimpar');
      const btnSalvarJust = document.getElementById('btnSalvarJust');
      const btnAtualizar = document.getElementById('btnAtualizar');
      const buscaInput = document.getElementById('buscaInput');

      const painel = document.getElementById('painel');
      const corpo = document.getElementById('corpo');
      const descModal = document.getElementById('descModal');
      const justInput = document.getElementById('justInput');

      const modalEl = document.getElementById('modalJust');
      const modal = bootstrap.Modal.getOrCreateInstance(modalEl);

      let chaveDesvio = "";
      let ultChave = "";
      let ultTime = 0;
      let processando = false;

      setInterval(() => {
        const d = new Date();
        document.getElementById('clock').innerText =
          d.toLocaleDateString('pt-BR') + ' ' + d.toLocaleTimeString('pt-BR');
      }, 1000);

      function focar() {
        inputElem.focus();
        inputElem.select();
      }

      function limpar() {
        inputElem.value = "";
        ultChave = "";
        focar();
      }

      function getBusca() {
        return (buscaInput.value || "").trim();
      }

      async function executarRequisicao(chave) {
        processando = true;
        painel.className = "alert alert-warning border text-center my-0 py-3 fw-bold";
        painel.innerText = "⏳ Processando...";

        try {
          const res = await fetch('/api/bipar', {
            method: 'POST',
            headers: {'Content-Type': 'application/json'},
            body: JSON.stringify({chave})
          });

          const data = await res.json();

          if (!res.ok || !data.sucesso) {
            painel.className = "alert alert-danger border text-center my-0 py-3 fw-bold";
            painel.innerText = "❌ " + (data.mensagem || "Falha no processamento.");
          } else if (data.tipo === "1ª Bipagem") {
            painel.className = "alert alert-primary border text-center my-0 py-3 fw-bold";
            painel.innerText = "✅ 1ª Bipagem Gravada: NF " + data.dados.numero_nf + " às " + data.dados.hora_bip1;
            inputElem.value = "";
          } else {
            painel.className = "alert alert-danger border text-center my-0 py-3 fw-bold";
            const prefixo = data.dados.na_base_comparacao ? "📚 Base histórica • " : "";
            painel.innerText = "🚨 " + prefixo + "ALERTA DE DESVIO: NF " + data.dados.numero_nf + " | Tempo: " + data.dados.tempo_decorrido;
            chaveDesvio = data.dados.chave;
            descModal.innerText = "A NF " + data.dados.numero_nf + " já possui saída anterior.\nTempo decorrido: " + data.dados.tempo_decorrido;
            justInput.value = "";
            btnSalvarJust.innerText = "Gravar Justificativa e Enviar E-mail";
            btnSalvarJust.disabled = false;
            modal.show();
          }
        } catch (e) {
          console.error(e);
          painel.className = "alert alert-danger border text-center my-0 py-3 fw-bold";
          painel.innerText = "❌ Falha de comunicação com o servidor.";
        } finally {
          processando = false;
          await carregar();
          focar();
        }
      }

      function biparManual() {
        if (processando) return;
        const chave = (inputElem.value || "").trim();
        if (!chave) {
          alert("Por favor, bip a nota fiscal ou cole a chave primeiro!");
          focar();
          return;
        }
        ultChave = "";
        executarRequisicao(chave);
      }

      function biparAutomatico() {
        if (processando) return;
        const chave = (inputElem.value || "").trim();
        if (!chave) return;

        const now = Date.now();
        if (chave === ultChave && (now - ultTime < 2500)) return; // anti-rebote
        ultChave = chave;
        ultTime = now;

        executarRequisicao(chave);
      }

      async function salvarJust() {
        const just = (justInput.value || "").trim();
        if (!just) {
          alert("Digite o motivo obrigatório.");
          return;
        }

        btnSalvarJust.disabled = true;
        btnSalvarJust.innerText = "Salvando e Enviando E-mail...";

        try {
          const res = await fetch('/api/justificar', {
            method: 'POST',
            headers: {'Content-Type': 'application/json'},
            body: JSON.stringify({chave: chaveDesvio, justificativa: just})
          });
          const data = await res.json();
          if (!res.ok || !data.sucesso) {
            alert(data.mensagem || "Falha ao salvar justificativa.");
          }
          inputElem.value = "";
        } catch (e) {
          console.error(e);
          alert("Falha ao salvar justificativa.");
        }

        modal.hide();
        await carregar();
        focar();
      }

      async function carregar() {
        try {
          const busca = encodeURIComponent(getBusca());
          const res = await fetch('/api/historico?busca=' + busca);
          const lista = await res.json();

          corpo.innerHTML = "";
          lista.forEach(n => {
            let st = "";
            const status = n.status || "";
            if (status.includes("DESVIO")) st = '<span class="badge bg-danger">🚨 DESVIO</span>';
            else if (status.includes("PENDENTE")) st = '<span class="badge bg-warning text-dark">⚠️ PENDENTE</span>';
            else st = '<span class="badge bg-success">REGULAR</span>';

            corpo.innerHTML += `<tr class="text-center">
              <td class="fw-bold">${n.numero_nf || '-'}</td>
              <td class="font-monospace small text-start">${n.chave || '-'}</td>
              <td>${n.serie || '-'}</td>
              <td>${n.data_bip1 || '-'}</td>
              <td>${n.hora_bip1 || '-'}</td>
              <td>${n.data_bip2 || '-'}</td>
              <td>${n.hora_bip2 || '-'}</td>
              <td class="fw-bold">${n.tempo_decorrido || '-'}</td>
              <td>${st}</td>
              <td class="text-start small">${n.justificativa || '-'}</td>
            </tr>`;
          });
        } catch (e) {
          console.error(e);
        }
      }

      btnRegistrar.addEventListener('click', biparManual);
      btnLimpar.addEventListener('click', limpar);
      btnSalvarJust.addEventListener('click', salvarJust);
      btnAtualizar.addEventListener('click', carregar);

      buscaInput.addEventListener('keydown', (e) => {
        if (e.key === 'Enter') {
          e.preventDefault();
          carregar();
        }
      });

      inputElem.addEventListener('keydown', (e) => {
        if (e.key === 'Enter') {
          e.preventDefault();
          biparAutomatico();
        }
      });

      inputElem.addEventListener('input', (e) => {
        const v = (e.target.value || '').replace(/\\D/g, '');
        if (v.length === 44 && v !== ultChave) {
          setTimeout(biparAutomatico, 200);
        }
      });

      carregar();
      focar();
    });
  </script>
</body>
</html>
"""


# =========================================================
# ROTAS
# =========================================================
@app.route("/")
def index():
    return render_template_string(HTML)


@app.route("/api/bipar", methods=["POST"])
def api_bipar():
    payload = request.get_json(silent=True) or {}
    chave_raw = (payload.get("chave") or "").strip()

    d = sanitizar_e_extrair_chave(chave_raw)
    if not d["valida"]:
        return jsonify({
            "sucesso": False,
            "mensagem": f"Chave inválida ({len(d['chave'])} dígitos). Precisa ter 44."
        }), 400

    agora = agora_local()
    dt_str = agora.strftime("%d/%m/%Y")
    hr_str = agora.strftime("%H:%M:%S")
    iso_str = agora.isoformat()

    conn = obter_conexao()
    c = conn.cursor()

    try:
        # Busca na base oculta
        c.execute("SELECT * FROM base_comparacao WHERE chave = ?", (d["chave"],))
        row_base = c.fetchone()
        existe_base_comp = row_base is not None

        # Busca no histórico operacional
        c.execute("SELECT * FROM notas WHERE chave = ?", (d["chave"],))
        nota = c.fetchone()

        # -----------------------------------------------------
        # Não existe em notas ainda
        # -----------------------------------------------------
        if not nota:
            # Se existe na base oculta -> já entra como DESVIO
            if existe_base_comp:
                dt_ref = parse_data_hora_ref(row_base["data_ref"], row_base["hora_ref"])
                if dt_ref is None:
                    dt_ref = agora

                data_ref = dt_ref.strftime("%d/%m/%Y")
                hora_ref = dt_ref.strftime("%H:%M:%S")
                tempo_txt, mins = calcular_diferenca(dt_ref, agora)

                c.execute("""
                    INSERT INTO notas (
                        chave, numero_nf, serie,
                        data_bip1, hora_bip1, dt_completa_bip1,
                        data_bip2, hora_bip2, dt_completa_bip2,
                        tempo_decorrido, minutos_decorridos,
                        status, justificativa, qtd_bipagens, observacao
                    ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """, (
                    d["chave"], d["numero_nf"], d["serie"],
                    data_ref, hora_ref, dt_ref.isoformat(),
                    dt_str, hr_str, iso_str,
                    tempo_txt, mins,
                    "DESVIO PENDENTE JUSTIFICATIVA", "", 2, "BASE_COMPARACAO"
                ))

                conn.commit()
                return jsonify({
                    "sucesso": True,
                    "tipo": "2ª Bipagem",
                    "dados": {
                        "chave": d["chave"],
                        "numero_nf": d["numero_nf"],
                        "serie": d["serie"],
                        "data_bip1": data_ref,
                        "hora_bip1": hora_ref,
                        "data_bip2": dt_str,
                        "hora_bip2": hr_str,
                        "tempo_decorrido": tempo_txt,
                        "na_base_comparacao": True
                    }
                }), 200

            # Fluxo regular
            c.execute("""
                INSERT INTO notas (
                    chave, numero_nf, serie,
                    data_bip1, hora_bip1, dt_completa_bip1,
                    status, justificativa, qtd_bipagens, observacao
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """, (
                d["chave"], d["numero_nf"], d["serie"],
                dt_str, hr_str, iso_str,
                "REGULAR", "", 1, None
            ))

            conn.commit()
            return jsonify({
                "sucesso": True,
                "tipo": "1ª Bipagem",
                "dados": {
                    "chave": d["chave"],
                    "numero_nf": d["numero_nf"],
                    "serie": d["serie"],
                    "data_bip1": dt_str,
                    "hora_bip1": hr_str,
                    "na_base_comparacao": False
                }
            }), 200

        # -----------------------------------------------------
        # Já existe em notas -> desvio normal (2ª+ leitura)
        # -----------------------------------------------------
        dt1 = parse_iso_db(nota["dt_completa_bip1"]) or agora
        tempo_txt, mins = calcular_diferenca(dt1, agora)

        try:
            qtd = int(nota["qtd_bipagens"] or 1) + 1
        except Exception:
            qtd = 2

        observacao = nota["observacao"] or ""
        if existe_base_comp and "BASE_COMPARACAO" not in observacao:
            observacao = "BASE_COMPARACAO"

        c.execute("""
            UPDATE notas
               SET data_bip2 = ?,
                   hora_bip2 = ?,
                   dt_completa_bip2 = ?,
                   tempo_decorrido = ?,
                   minutos_decorridos = ?,
                   status = ?,
                   qtd_bipagens = ?,
                   observacao = ?
             WHERE chave = ?
        """, (
            dt_str, hr_str, iso_str,
            tempo_txt, mins,
            "DESVIO PENDENTE JUSTIFICATIVA",
            qtd,
            observacao if observacao else None,
            d["chave"]
        ))

        conn.commit()
        return jsonify({
            "sucesso": True,
            "tipo": "2ª Bipagem",
            "dados": {
                "chave": d["chave"],
                "numero_nf": d["numero_nf"],
                "serie": d["serie"],
                "data_bip1": nota["data_bip1"],
                "hora_bip1": nota["hora_bip1"],
                "data_bip2": dt_str,
                "hora_bip2": hr_str,
                "tempo_decorrido": tempo_txt,
                "na_base_comparacao": existe_base_comp
            }
        }), 200

    except Exception as e:
        conn.rollback()
        return jsonify({"sucesso": False, "mensagem": f"Erro ao processar bipagem: {str(e)}"}), 500
    finally:
        conn.close()


@app.route("/api/justificar", methods=["POST"])
def api_justificar():
    req = request.get_json(silent=True) or {}
    chave = (req.get("chave") or "").strip()
    justificativa = (req.get("justificativa") or "").strip()

    if not chave or not justificativa:
        return jsonify({"sucesso": False, "mensagem": "Chave e justificativa são obrigatórios."}), 400

    conn = obter_conexao()
    c = conn.cursor()

    try:
        c.execute("UPDATE notas SET justificativa = ?, status = 'DESVIO REGISTRADO' WHERE chave = ?", (justificativa, chave))
        c.execute("SELECT * FROM notas WHERE chave = ?", (chave,))
        row = c.fetchone()
        conn.commit()
    except Exception as e:
        conn.rollback()
        conn.close()
        return jsonify({"sucesso": False, "mensagem": f"Falha ao salvar justificativa: {e}"}), 500
    finally:
        conn.close()

    if row:
        def disparar_email():
            cfg = carregar_config()
            destinatarios = cfg.get("destinatarios", [])
            if destinatarios:
                payload = {
                    "chave": chave,
                    "numero_nf": row["numero_nf"],
                    "serie": row["serie"],
                    "data_bip1": row["data_bip1"],
                    "hora_bip1": row["hora_bip1"],
                    "data_bip2": row["data_bip2"],
                    "hora_bip2": row["hora_bip2"],
                    "tempo_decorrido": row["tempo_decorrido"],
                    "justificativa": justificativa,
                }
                assunto = f"[ALERTA DE DESVIO - JUSTIFICADO] NF {row['numero_nf']}"
                html = template_email_desvio(payload)
                enviar_email_smtp(assunto, html, destinatarios)

        threading.Thread(target=disparar_email, daemon=True).start()

    return jsonify({"sucesso": True}), 200


@app.route("/api/historico")
def api_historico():
    busca = (request.args.get("busca") or "").strip()

    conn = obter_conexao()
    c = conn.cursor()

    if busca:
        like = f"%{busca}%"
        c.execute("""
            SELECT *
            FROM notas
            WHERE chave LIKE ?
               OR numero_nf LIKE ?
               OR serie LIKE ?
               OR IFNULL(status, '') LIKE ?
               OR IFNULL(justificativa, '') LIKE ?
               OR IFNULL(data_bip1, '') LIKE ?
               OR IFNULL(data_bip2, '') LIKE ?
            ORDER BY dt_completa_bip1 DESC
        """, (like, like, like, like, like, like, like))
    else:
        c.execute("SELECT * FROM notas ORDER BY dt_completa_bip1 DESC")

    rows = c.fetchall()
    conn.close()
    return jsonify([dict(r) for r in rows]), 200


@app.route("/api/exportar")
def api_exportar():
    conn = obter_conexao()
    registros = conn.execute("SELECT * FROM notas ORDER BY dt_completa_bip1 DESC").fetchall()
    conn.close()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Portaria"

    headers = [
        "Chave de Acesso", "Número NF", "Série", "Data 1ª Bip", "Hora 1ª Bip",
        "Data 2ª Bip", "Hora 2ª Bip", "Diferença", "Status", "Justificativa", "Observação"
    ]
    ws.append(headers)

    for r in registros:
        ws.append([
            r["chave"], r["numero_nf"], r["serie"],
            r["data_bip1"], r["hora_bip1"],
            r["data_bip2"] or "-", r["hora_bip2"] or "-",
            r["tempo_decorrido"] or "-", r["status"] or "-", r["justificativa"] or "-",
            r["observacao"] or "-"
        ])

    stream = BytesIO()
    wb.save(stream)
    stream.seek(0)

    nome = f"Relatorio_Portaria_{agora_local().strftime('%Y%m%d_%H%M%S')}.xlsx"
    return send_file(
        stream,
        as_attachment=True,
        download_name=nome,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


@app.route("/api/importar-base-comparacao", methods=["POST"])
def api_importar_base_comparacao():
    """
    multipart/form-data:
      - arquivo: .xlsx (obrigatório)
      - aba: nome da aba (opcional)
    """
    if "arquivo" not in request.files:
        return jsonify({"sucesso": False, "mensagem": "Envie o arquivo no campo 'arquivo'."}), 400

    file = request.files["arquivo"]
    nome_arquivo = (file.filename or "arquivo.xlsx").strip()
    aba_req = (request.form.get("aba") or "").strip()

    if not nome_arquivo.lower().endswith(".xlsx"):
        return jsonify({"sucesso": False, "mensagem": "Formato inválido. Envie .xlsx"}), 400

    try:
        wb = openpyxl.load_workbook(file, data_only=True)
    except Exception as e:
        return jsonify({"sucesso": False, "mensagem": f"Falha ao ler Excel: {e}"}), 400

    if aba_req:
        if aba_req not in wb.sheetnames:
            return jsonify({
                "sucesso": False,
                "mensagem": f"Aba '{aba_req}' não encontrada. Abas disponíveis: {', '.join(wb.sheetnames)}"
            }), 400
        ws = wb[aba_req]
    else:
        ws = wb[wb.sheetnames[0]]

    colmap = detectar_colunas_base(ws)
    if not colmap["chave"]:
        return jsonify({
            "sucesso": False,
            "mensagem": "Não encontrei coluna de CHAVE. Use cabeçalhos como 'Chave' ou 'Chave de Acesso'."
        }), 400

    validos, rejeitados = extrair_registros_base(ws, colmap)

    if not validos:
        return jsonify({
            "sucesso": False,
            "mensagem": "Nenhum registro válido (44 dígitos) encontrado para importar.",
            "rejeitados": len(rejeitados)
        }), 400

    agora_iso = agora_local().isoformat()

    conn = obter_conexao()
    c = conn.cursor()

    inseridos = 0
    atualizados = 0

    try:
        for row in validos:
            c.execute("SELECT 1 FROM base_comparacao WHERE chave = ?", (row["chave"],))
            exists = c.fetchone() is not None

            c.execute("""
                INSERT INTO base_comparacao
                    (chave, numero_nf, serie, data_ref, hora_ref, origem_arquivo, importado_em)
                VALUES (?, ?, ?, ?, ?, ?, ?)
                ON CONFLICT(chave) DO UPDATE SET
                    numero_nf=excluded.numero_nf,
                    serie=excluded.serie,
                    data_ref=excluded.data_ref,
                    hora_ref=excluded.hora_ref,
                    origem_arquivo=excluded.origem_arquivo,
                    importado_em=excluded.importado_em
            """, (
                row["chave"], row["numero_nf"], row["serie"],
                row["data_ref"], row["hora_ref"], nome_arquivo, agora_iso
            ))

            if exists:
                atualizados += 1
            else:
                inseridos += 1

        conn.commit()
    except Exception as e:
        conn.rollback()
        conn.close()
        return jsonify({"sucesso": False, "mensagem": f"Falha na importação: {e}"}), 500
    finally:
        conn.close()

    return jsonify({
        "sucesso": True,
        "mensagem": "Base de comparação importada com sucesso.",
        "arquivo": nome_arquivo,
        "aba": ws.title,
        "linhas_validas": len(validos),
        "inseridos": inseridos,
        "atualizados": atualizados,
        "rejeitados": len(rejeitados),
        "exemplos_rejeitados": rejeitados[:10],
        "colunas_detectadas": colmap
    }), 200


@app.route("/api/base-comparacao/resumo")
def api_base_comparacao_resumo():
    conn = obter_conexao()
    c = conn.cursor()

    c.execute("SELECT COUNT(*) AS total FROM base_comparacao")
    total = c.fetchone()["total"]

    c.execute("""
        SELECT origem_arquivo, importado_em
        FROM base_comparacao
        ORDER BY importado_em DESC
        LIMIT 1
    """)
    last = c.fetchone()

    conn.close()
    return jsonify({
        "sucesso": True,
        "total_registros_base": total,
        "ultimo_arquivo": (last["origem_arquivo"] if last else None),
        "ultimo_importado_em": (last["importado_em"] if last else None),
    }), 200


@app.route("/api/base-comparacao/limpar", methods=["POST"])
def api_base_comparacao_limpar():
    conn = obter_conexao()
    c = conn.cursor()
    try:
        c.execute("DELETE FROM base_comparacao")
        conn.commit()
    except Exception as e:
        conn.rollback()
        conn.close()
        return jsonify({"sucesso": False, "mensagem": f"Erro ao limpar base: {e}"}), 500
    finally:
        conn.close()

    return jsonify({"sucesso": True, "mensagem": "Base de comparação limpa com sucesso."}), 200


# =========================================================
# INIT
# =========================================================
inicializar_banco()

if __name__ == "__main__":
    port = int(os.environ.get("PORT", 5000))
    app.run(host="0.0.0.0", port=port)
